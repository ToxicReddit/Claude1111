"""Build the film watch-guide workbook from the curated catalog + a Letterboxd export.

Usage:
    python movies/build_watch_guide.py <letterboxd-export.zip | unzipped-folder> [out.xlsx]

Letterboxd is used only to mark what you've already seen. The "Seen" columns are
live COUNTIF formulas against the 'Letterboxd Watched' sheet, so pasting a newer
watched.csv into that sheet updates every tab.
"""
import csv
import io
import re
import sys
import unicodedata
import zipfile
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from catalog import NICHES  # noqa: E402

FONT = "Arial"
MAXR = 3000  # formula ranges on the Letterboxd sheet leave room for a bigger export

CAT_COLORS = {
    "01": "DCE6F1", "02": "E4DFEC", "03": "DDEBF7", "04": "FCE4D6", "05": "E2EFDA",
    "06": "FFF2CC", "07": "EDEDED", "08": "F8CBAD", "09": "D9E1F2", "10": "C6E0B4",
}
HDR_FILL = PatternFill("solid", fgColor="1F2A44")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=16, bold=True, color="1F2A44")
SUB = Font(name=FONT, size=11, bold=True, color="1F2A44")
MUTED = Font(name=FONT, size=9, italic=True, color="595959")
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(bottom=THIN)
SEEN_FILL = PatternFill("solid", fgColor="E2F0D9")
NEXT_FILL = PatternFill("solid", fgColor="FFE699")
TIER_NAMES = {1: "1 · Start Here", 2: "2 · Core", 3: "3 · Deep Cut"}
TIER_WEIGHT = {1: 3, 2: 2, 3: 1}


# ---------------------------------------------------------------- Letterboxd
def read_export(src):
    """Return {relative_path: text} for every CSV in the export (zip or folder)."""
    src = Path(src)
    out = {}
    if src.is_file() and src.suffix == ".zip":
        with zipfile.ZipFile(src) as z:
            for n in z.namelist():
                if n.endswith(".csv"):
                    out[n] = z.read(n).decode("utf-8-sig")
    else:
        for p in src.rglob("*.csv"):
            out[str(p.relative_to(src))] = p.read_text(encoding="utf-8-sig")
    return out


def rows_from_list_csv(text):
    """Letterboxd list CSVs have a metadata block before the 'Position,...' header."""
    lines = text.splitlines()
    for i, line in enumerate(lines):
        if line.startswith("Position,"):
            return list(csv.DictReader(io.StringIO("\n".join(lines[i:]))))
    return []


def norm(s):
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace("&", " and ").replace("½", " 1/2")
    s = re.sub(r"[^a-z0-9/ ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_noart(s):
    return re.sub(r"^(the|a|an|le|la|les|l|il|el|der|die|das) ", "", norm(s))


# Catalog title -> Letterboxd title, where the two differ beyond punctuation/accents.
ALIASES = {
    ("Ringu", 1998): "Ring",
    ("Ju-on: The Grudge", 2002): "Ju-on: The Grudge",
    ("Wallace & Gromit: The Wrong Trousers", 1993): "The Wrong Trousers",
    ("October (Ten Days that Shook the World)", 1928): "October",
    ("Vidas Secas", 1963): "Barren Lives",
    ("Horror of Dracula", 1958): "Dracula",
    ("Zombie Flesh Eaters", 1979): "Zombie",
    ("Life of Brian", 1979): "Monty Python's Life of Brian",
    ("Etre et avoir", 2002): "To Be and to Have",
    ("7 Up!", 1964): "Seven Up!",
    ("Lone Wolf and Cub: Sword of Vengeance", 1972): "Lone Wolf and Cub: Sword of Vengeance",
}


def build_lb_index(files):
    """Index of every film Letterboxd names anywhere in the export (watched, lists…)."""
    idx = defaultdict(list)  # normalized title -> [(name, year, uri)]
    seen_uri = set()

    def add(name, year, uri):
        try:
            year = int(float(year))
        except (TypeError, ValueError):
            return
        key = (name, year)
        if key in seen_uri:
            return
        seen_uri.add(key)
        for k in {norm(name), norm_noart(name)}:
            idx[k].append((name, year, uri))

    for path, text in files.items():
        base = path.split("/")[-1]
        if base in ("watched.csv", "watchlist.csv", "ratings.csv", "diary.csv"):
            for r in csv.DictReader(io.StringIO(text)):
                add(r.get("Name", ""), r.get("Year"), r.get("Letterboxd URI", ""))
        elif "lists/" in path:
            for r in rows_from_list_csv(text):
                add(r.get("Name", ""), r.get("Year"), r.get("URL", ""))
    return idx


def canonical(title, year, idx):
    """Map a catalog film to its Letterboxd name/year/URI when the export knows it."""
    lookup = ALIASES.get((title, year), title)
    for k in (norm(lookup), norm_noart(lookup)):
        cands = idx.get(k, [])
        for dy in (0, 1, -1, 2, -2):
            for name, y, uri in cands:
                if y == year + dy:
                    return name, y, uri
    return title, year, ""


# ---------------------------------------------------------------- helpers
def style_header(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def lb_link(title, year, uri):
    if uri:
        return uri
    q = re.sub(r"\s+", "+", re.sub(r"[^\w\s]", "", title).strip())
    return f"https://letterboxd.com/search/films/{q}+{year}/"


# ---------------------------------------------------------------- main
def main(src, out):
    files = read_export(src)
    watched = []
    for path, text in files.items():
        if path.split("/")[-1] == "watched.csv" and "/" not in path.strip("/"):
            watched = list(csv.DictReader(io.StringIO(text)))
    if not watched:  # fall back to any watched.csv
        for path, text in files.items():
            if path.endswith("watched.csv"):
                watched = list(csv.DictReader(io.StringIO(text)))
    idx = build_lb_index(files)

    # ---- flatten catalog into path rows + master films
    films = OrderedDict()   # (lbname, lbyear) -> info
    path_rows = []
    niche_rows = []
    warnings = []
    for n_i, n in enumerate(NICHES, 1):
        nid = f"N{n_i:03d}"
        seen_in_niche = set()
        step = 0
        for title, year, director, tier in n["films"]:
            name, y, uri = canonical(title, year, idx)
            key = (name, y)
            if key in seen_in_niche:
                warnings.append(f"duplicate in niche {n['niche']!r}: {title} ({year})")
                continue
            seen_in_niche.add(key)
            step += 1
            f = films.get(key)
            if f is None:
                f = films[key] = dict(name=name, year=y, director=director, uri=uri,
                                      cats=OrderedDict(), niches=[], tiers=[], score=0)
            elif f["director"] != director and director.split(",")[0] not in f["director"]:
                warnings.append(f"director mismatch for {name} ({y}): {f['director']!r} vs {director!r}")
            f["cats"][n["category"]] = True
            f["niches"].append(n["niche"])
            f["tiers"].append(tier)
            f["score"] += TIER_WEIGHT[tier]
            path_rows.append(dict(nid=nid, cat=n["category"], sub=n["subcategory"], niche=n["niche"],
                                  step=step, tier=tier, name=name, year=y, director=f["director"]))
        niche_rows.append(dict(nid=nid, cat=n["category"], sub=n["subcategory"], niche=n["niche"],
                               sources=n["sources"], note=n["note"], count=step))

    categories = list(OrderedDict((n["category"], 1) for n in NICHES))
    watched_keys = {(r["Name"].lower(), int(float(r["Year"]))) for r in watched if r.get("Year")}
    seen_count = sum((f["name"].lower(), f["year"]) in watched_keys for f in films.values())

    wb = Workbook()

    # ================================================================ README
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    set_widths(ws, [3, 30, 95])
    ws["B2"] = "The Film Watch Guide"
    ws["B2"].font = TITLE
    ws["B3"] = ("An ordered map of cinema by category, subcategory, and niche. It draws on film-studies textbooks, "
                "critics' polls, labels/curators (Criterion, Arrow, BFI, Masters of Cinema, Second Run, the World Cinema Project), "
                "and your own Letterboxd history, which is used only to mark what you've seen.")
    ws["B3"].font = BODY
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B3:C3")
    ws.row_dimensions[3].height = 45

    r = 5
    ws.cell(row=r, column=2, value="At a glance").font = SUB
    stats = [
        ("Unique films", "=COUNTA('Master List'!A2:A%d)" % (len(films) + 1)),
        ("Seen (from Letterboxd)", "=SUM('Master List'!G2:G%d)" % (len(films) + 1)),
        ("% seen", "=IF(C6=0,0,C7/C6)"),
        ("Watch-path entries (a film can sit in many buckets)", "=COUNTA('Watch Paths'!G2:G%d)" % (len(path_rows) + 1)),
        ("Categories / subcategories / niches", f"{len(categories)} / {len({(n['category'], n['subcategory']) for n in NICHES})} / {len(NICHES)}"),
    ]
    for i, (label, val) in enumerate(stats, r + 1):
        ws.cell(row=i, column=2, value=label).font = BODY
        c = ws.cell(row=i, column=3, value=val)
        c.font = BOLD
        c.alignment = Alignment(horizontal="left")
    ws["C8"].number_format = "0.0%"

    r = 12
    ws.cell(row=r, column=2, value="How the tabs work").font = SUB
    guide = [
        ("Dashboard", "Progress by category and subcategory. The bars fill in as you log films on Letterboxd."),
        ("Priority Queue", "Every film ranked by how central it is across the whole map (weighted bucket count). "
                           "Filter Seen = blank, then work down the list for the best overall order."),
        ("Watch Paths", "The main list: every category → subcategory → niche in recommended viewing order (Step). "
                        "'▶ NEXT' marks the first unseen film in each niche. Filter by Category, Niche, or Tier."),
        ("Taxonomy", "One row per niche: its academic/critical sources, why it's grouped, your progress, and your next film."),
        ("Master List", "One row per unique film, with every bucket it belongs to and a Letterboxd link."),
        ("Bibliography", "The books, essays, polls, and label catalogs behind the groupings."),
        ("Letterboxd Watched", "Your watched.csv. To refresh, paste a newer export's Date/Name/Year/URI into columns A–D. "
                               "Column E rebuilds the match key, and every Seen mark updates."),
    ]
    for i, (a, b) in enumerate(guide, r + 1):
        ws.cell(row=i, column=2, value=a).font = BOLD
        c = ws.cell(row=i, column=3, value=b)
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30

    r = 21
    ws.cell(row=r, column=2, value="How the ordering works").font = SUB
    order = [
        ("Tier 1 · Start Here", "The gateway films for a niche: canonical, accessible, and they teach you how to watch the rest."),
        ("Tier 2 · Core", "The essential body of the niche, mostly in chronological or influence order, so you can see the style develop."),
        ("Tier 3 · Deep Cut", "Rarer, harder, or more specialized titles for once the niche clicks."),
        ("Step", "The recommended order within a niche. Movement/era niches are chronological. Genre and auteur niches go gateway → chronology → deep cuts."),
        ("Multiple buckets", "Films deliberately appear in many niches (Vertigo sits in history, Hitchcock, theory, canons…). "
                             "The Priority Queue turns that overlap into a centrality score."),
        ("Suggested route", "(1) Foundations → (2) History Spine, alongside (5) Auteurs and (4) Genres by mood → "
                            "(3) World Cinemas → (6) Theory lenses once you want a vocabulary. Dip into (7) Themes, (8) Niches, "
                            "(9) Canons, and (10) Bridges whenever you like."),
        ("Legend", "✓ = logged on your Letterboxd. ▶ NEXT = first unseen film in that niche. Green rows = seen."),
        ("Source note", "Poll and list selections (S&S 2022, AFI 2007, BFI 100, Ebert) are subsets shown in viewing order, not rank. "
                        "Label entries (Arrow, BFI Flipside, Masters of Cinema, Second Run) describe each label's catalog focus; "
                        "editions go in and out of print, so check the label for current releases. "
                        "'Composite' curricula are built from standard textbooks, not any single university's syllabus."),
    ]
    for i, (a, b) in enumerate(order, r + 1):
        ws.cell(row=i, column=2, value=a).font = BOLD
        c = ws.cell(row=i, column=3, value=b)
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 44 if len(b) > 150 else 30

    # ================================================================ Letterboxd Watched
    wl = wb.create_sheet("Letterboxd Watched")
    heads = ["Date", "Name", "Year", "Letterboxd URI", "Match Key (formula)"]
    wl.append(heads)
    style_header(wl, 1, len(heads))
    for i, rrow in enumerate(watched, 2):
        wl.cell(row=i, column=1, value=rrow.get("Date"))
        wl.cell(row=i, column=2, value=rrow.get("Name"))
        try:
            wl.cell(row=i, column=3, value=int(float(rrow.get("Year"))))
        except (TypeError, ValueError):
            wl.cell(row=i, column=3, value=None)
        wl.cell(row=i, column=4, value=rrow.get("Letterboxd URI"))
    for i in range(2, MAXR + 1):
        wl.cell(row=i, column=5, value=f'=IF(B{i}="","",LOWER(B{i})&"|"&C{i})')
    for row in wl.iter_rows(min_row=2, max_row=len(watched) + 1, max_col=5):
        for c in row:
            c.font = BODY
    set_widths(wl, [12, 44, 8, 30, 50])
    wl.freeze_panes = "A2"
    SEEN_RANGE = f"'Letterboxd Watched'!$E$2:$E${MAXR}"

    # ================================================================ Master List
    wm = wb.create_sheet("Master List")
    cat_short = [c.split(" ", 1)[0] for c in categories]
    heads = (["Title", "Year", "Decade", "Director", "Seen", "Buckets", "Seen #", "Centrality Score",
              "Best Tier", "Primary Category", "All Niches", "Letterboxd", "Match Key"]
             + [f"In Cat {s}" for s in cat_short])
    wm.append(heads)
    style_header(wm, 1, len(heads), 42)
    master_order = sorted(films.values(), key=lambda f: (norm_noart(f["name"]), f["year"]))
    for i, f in enumerate(master_order, 2):
        f["mrow"] = i - 1  # position inside the Master List data range
        vals = [f["name"], f["year"], f"=FLOOR(B{i},10)&\"s\"", f["director"],
                f'=IF(G{i}=1,"✓","")', len(f["niches"]), f"=IF(COUNTIF({SEEN_RANGE},M{i})>0,1,0)",
                f["score"], TIER_NAMES[min(f["tiers"])], next(iter(f["cats"])), "; ".join(f["niches"]),
                f'=HYPERLINK("{lb_link(f["name"], f["year"], f["uri"])}","open")',
                f'=LOWER(A{i})&"|"&B{i}']
        vals += [1 if c in f["cats"] else 0 for c in categories]
        for j, v in enumerate(vals, 1):
            cell = wm.cell(row=i, column=j, value=v)
            cell.font = LINK if j == 12 else BODY
        wm.cell(row=i, column=5).alignment = Alignment(horizontal="center")
    last_m = len(master_order) + 1
    set_widths(wm, [42, 7, 8, 30, 6, 8, 7, 10, 14, 36, 80, 10, 40] + [8] * len(categories))
    wm.freeze_panes = "B2"
    wm.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{last_m}"
    wm.conditional_formatting.add(f"A2:L{last_m}", FormulaRule(formula=[f"$G2=1"], fill=SEEN_FILL))
    for col in ("G", "M"):
        wm.column_dimensions[col].hidden = True
    for k in range(14, 14 + len(categories)):
        wm.column_dimensions[get_column_letter(k)].hidden = True

    # ================================================================ Watch Paths
    wp = wb.create_sheet("Watch Paths")
    heads = ["Niche ID", "Category", "Subcategory", "Niche", "Step", "Tier", "Title", "Year", "Director",
             "Seen", "Next Up", "Buckets", "Seen #", "Next ID", "Master Row", "Unseen Before"]
    wp.append(heads)
    style_header(wp, 1, len(heads))
    last_p = len(path_rows) + 1
    for i, pr in enumerate(path_rows, 2):
        vals = [pr["nid"], pr["cat"], pr["sub"], pr["niche"], pr["step"], TIER_NAMES[pr["tier"]],
                pr["name"], pr["year"], pr["director"],
                f'=IF(M{i}=1,"✓","")',
                f'=IF(AND(M{i}=0,P{i}=0),"▶ NEXT","")',
                f"=INDEX('Master List'!$F$2:$F${last_m},O{i})",
                f"=INDEX('Master List'!$G$2:$G${last_m},O{i})",
                f'=IF(K{i}="","",A{i})',
                films[(pr["name"], pr["year"])]["mrow"],
                0 if i == 2 else f"=IF(A{i}=A{i - 1},P{i - 1}+1-M{i - 1},0)"]
        fill = PatternFill("solid", fgColor=CAT_COLORS[pr["cat"][:2]])
        for j, v in enumerate(vals, 1):
            cell = wp.cell(row=i, column=j, value=v)
            cell.font = BOLD if (j == 7 and pr["tier"] == 1) else BODY
            if j <= 4:
                cell.fill = fill
        wp.cell(row=i, column=10).alignment = Alignment(horizontal="center")
    set_widths(wp, [8, 30, 34, 44, 6, 14, 44, 7, 30, 6, 10, 8, 7, 8, 8, 8])
    wp.freeze_panes = "G2"
    wp.auto_filter.ref = f"A1:L{last_p}"
    wp.conditional_formatting.add(f"E2:L{last_p}", FormulaRule(formula=["$M2=1"], fill=SEEN_FILL))
    wp.conditional_formatting.add(f"K2:K{last_p}", CellIsRule(operator="equal", formula=['"▶ NEXT"'],
                                                              fill=NEXT_FILL, font=Font(name=FONT, bold=True)))
    for col in ("M", "N", "O", "P"):
        wp.column_dimensions[col].hidden = True

    # ================================================================ Priority Queue
    wq = wb.create_sheet("Priority Queue")
    heads = ["Rank", "Title", "Year", "Director", "Seen", "Centrality Score", "Buckets", "Best Tier",
             "Primary Category", "Seen #", "Master Row"]
    wq.append(heads)
    style_header(wq, 1, len(heads))
    q = sorted(films.values(), key=lambda f: (-f["score"], -len(f["cats"]), f["year"]))
    for i, f in enumerate(q, 2):
        vals = [i - 1, f["name"], f["year"], f["director"], f'=IF(J{i}=1,"✓","")', f["score"],
                len(f["niches"]), TIER_NAMES[min(f["tiers"])], next(iter(f["cats"])),
                f"=INDEX('Master List'!$G$2:$G${last_m},K{i})", f["mrow"]]
        for j, v in enumerate(vals, 1):
            wq.cell(row=i, column=j, value=v).font = BODY
        wq.cell(row=i, column=5).alignment = Alignment(horizontal="center")
    last_q = len(q) + 1
    set_widths(wq, [6, 44, 7, 32, 6, 10, 8, 14, 38, 7, 8])
    wq.freeze_panes = "C2"
    wq.auto_filter.ref = f"A1:I{last_q}"
    wq.conditional_formatting.add(f"A2:I{last_q}", FormulaRule(formula=["$J2=1"], fill=SEEN_FILL))
    wq.conditional_formatting.add(f"F2:F{last_q}", DataBarRule(start_type="min", end_type="max", color="8EA9DB"))
    wq.column_dimensions["J"].hidden = True
    wq.column_dimensions["K"].hidden = True
    wq.cell(row=1, column=13, value="Score = Σ over every bucket (Start Here = 3, Core = 2, Deep Cut = 1). "
                                    "Higher means more central to the whole map. Filter Seen = blank and go top-down.").font = MUTED

    # ================================================================ Taxonomy
    wt = wb.create_sheet("Taxonomy")
    heads = ["Niche ID", "Category", "Subcategory", "Niche", "Films", "Seen", "% Seen", "Next Up",
             "Why this grouping / how to watch it", "Academic & critical sources"]
    wt.append(heads)
    style_header(wt, 1, len(heads))
    for i, nr in enumerate(niche_rows, 2):
        vals = [nr["nid"], nr["cat"], nr["sub"], nr["niche"],
                f"=COUNTIF('Watch Paths'!$A$2:$A${last_p},A{i})",
                f"=COUNTIFS('Watch Paths'!$A$2:$A${last_p},A{i},'Watch Paths'!$M$2:$M${last_p},1)",
                f"=IF(E{i}=0,0,F{i}/E{i})",
                f"=IFERROR(INDEX('Watch Paths'!$G$2:$G${last_p},MATCH(A{i},'Watch Paths'!$N$2:$N${last_p},0)),\"✓ complete\")",
                nr["note"], nr["sources"].replace("*", "")]
        fill = PatternFill("solid", fgColor=CAT_COLORS[nr["cat"][:2]])
        for j, v in enumerate(vals, 1):
            cell = wt.cell(row=i, column=j, value=v)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=j >= 9, vertical="top")
            if j <= 4:
                cell.fill = fill
        wt.cell(row=i, column=7).number_format = "0%"
    last_t = len(niche_rows) + 1
    set_widths(wt, [8, 28, 30, 42, 7, 7, 8, 36, 60, 70])
    wt.freeze_panes = "E2"
    wt.auto_filter.ref = f"A1:J{last_t}"
    wt.conditional_formatting.add(f"G2:G{last_t}", DataBarRule(start_type="num", start_value=0,
                                                               end_type="num", end_value=1, color="70AD47"))

    # ================================================================ Dashboard
    wd = wb.create_sheet("Dashboard", 1)
    wd.sheet_view.showGridLines = False
    wd["B2"] = "Progress Dashboard"
    wd["B2"].font = TITLE
    wd["B3"] = "Unique films per category. A film counts once per category even if it's in several of its niches."
    wd["B3"].font = MUTED
    heads = ["Category", "Unique Films", "Seen", "% Seen", "Niches", "Niches Complete"]
    for j, h in enumerate(heads, 2):
        wd.cell(row=5, column=j, value=h)
    for c in range(2, 8):
        cell = wd.cell(row=5, column=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
    for k, cat in enumerate(categories):
        i = 6 + k
        col = get_column_letter(14 + k)
        wd.cell(row=i, column=2, value=cat)
        wd.cell(row=i, column=3, value=f"=COUNTIF('Master List'!${col}$2:${col}${last_m},1)")
        wd.cell(row=i, column=4, value=f"=COUNTIFS('Master List'!${col}$2:${col}${last_m},1,'Master List'!$G$2:$G${last_m},1)")
        wd.cell(row=i, column=5, value=f"=IF(C{i}=0,0,D{i}/C{i})")
        wd.cell(row=i, column=6, value=f"=COUNTIF(Taxonomy!$B$2:$B${last_t},B{i})")
        wd.cell(row=i, column=7, value=f"=COUNTIFS(Taxonomy!$B$2:$B${last_t},B{i},Taxonomy!$G$2:$G${last_t},1)")
        for c in range(2, 8):
            cell = wd.cell(row=i, column=c)
            cell.font = BODY
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=CAT_COLORS[cat[:2]]) if c == 2 else PatternFill()
        wd.cell(row=i, column=5).number_format = "0%"
    end_c = 5 + len(categories)
    tr = end_c + 1
    wd.cell(row=tr, column=2, value="All films (unique)").font = BOLD
    wd.cell(row=tr, column=3, value=f"=COUNTA('Master List'!A2:A{last_m})").font = BOLD
    wd.cell(row=tr, column=4, value=f"=SUM('Master List'!G2:G{last_m})").font = BOLD
    wd.cell(row=tr, column=5, value=f"=IF(C{tr}=0,0,D{tr}/C{tr})").font = BOLD
    wd.cell(row=tr, column=5).number_format = "0%"
    wd.cell(row=tr, column=6, value=f"=SUM(F6:F{end_c})").font = BOLD
    wd.cell(row=tr, column=7, value=f"=SUM(G6:G{end_c})").font = BOLD
    wd.conditional_formatting.add(f"E6:E{end_c}", DataBarRule(start_type="num", start_value=0,
                                                             end_type="num", end_value=1, color="5B9BD5"))

    sr = tr + 3
    wd.cell(row=sr - 1, column=2, value="By subcategory (watch-path entries)").font = SUB
    heads = ["Category", "Subcategory", "Entries", "Seen", "% Seen", "Niches"]
    for j, h in enumerate(heads, 2):
        cell = wd.cell(row=sr, column=j, value=h)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
    subs = list(OrderedDict(((n["category"], n["subcategory"]), 1) for n in NICHES))
    for k, (cat, sub) in enumerate(subs):
        i = sr + 1 + k
        wd.cell(row=i, column=2, value=cat)
        wd.cell(row=i, column=3, value=sub)
        wd.cell(row=i, column=4, value=f"=COUNTIFS('Watch Paths'!$B$2:$B${last_p},B{i},'Watch Paths'!$C$2:$C${last_p},C{i})")
        wd.cell(row=i, column=5, value=f"=COUNTIFS('Watch Paths'!$B$2:$B${last_p},B{i},'Watch Paths'!$C$2:$C${last_p},C{i},'Watch Paths'!$M$2:$M${last_p},1)")
        wd.cell(row=i, column=6, value=f"=IF(D{i}=0,0,E{i}/D{i})")
        wd.cell(row=i, column=7, value=f"=COUNTIFS(Taxonomy!$B$2:$B${last_t},B{i},Taxonomy!$C$2:$C${last_t},C{i})")
        for c in range(2, 8):
            cell = wd.cell(row=i, column=c)
            cell.font = BODY
            cell.border = BOX
        wd.cell(row=i, column=2).fill = PatternFill("solid", fgColor=CAT_COLORS[cat[:2]])
        wd.cell(row=i, column=6).number_format = "0%"
    end_s = sr + len(subs)
    wd.conditional_formatting.add(f"F{sr + 1}:F{end_s}", DataBarRule(start_type="num", start_value=0,
                                                                    end_type="num", end_value=1, color="70AD47"))
    set_widths(wd, [3, 40, 40, 10, 10, 10, 16])

    # ================================================================ Bibliography
    wbib = wb.create_sheet("Bibliography")
    heads = ["Source", "Used for (niches)", "# Niches"]
    wbib.append(heads)
    style_header(wbib, 1, len(heads))
    src_map = OrderedDict()
    for n in NICHES:
        for s in re.split(r";\s+(?![^()]*\))", n["sources"]):
            s = s.replace("*", "").strip().rstrip(".")
            if s:
                src_map.setdefault(s, []).append(n["niche"])
    for i, (s, ns) in enumerate(sorted(src_map.items(), key=lambda kv: kv[0].lower()), 2):
        wbib.cell(row=i, column=1, value=s).font = BODY
        wbib.cell(row=i, column=2, value="; ".join(ns)).font = BODY
        wbib.cell(row=i, column=3, value=len(ns)).font = BODY
        wbib.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        wbib.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(wbib, [80, 90, 9])
    wbib.freeze_panes = "A2"
    wbib.auto_filter.ref = f"A1:C{len(src_map) + 1}"

    wb.move_sheet("Priority Queue", offset=-(wb.sheetnames.index("Priority Queue") - 2))
    wb.move_sheet("Watch Paths", offset=-(wb.sheetnames.index("Watch Paths") - 3))
    wb.move_sheet("Taxonomy", offset=-(wb.sheetnames.index("Taxonomy") - 4))
    wb.move_sheet("Master List", offset=-(wb.sheetnames.index("Master List") - 5))
    wb.move_sheet("Bibliography", offset=-(wb.sheetnames.index("Bibliography") - 6))
    for s in wb.worksheets:
        s.sheet_properties.tabColor = {"Start Here": "1F2A44", "Dashboard": "5B9BD5", "Priority Queue": "FFC000",
                                       "Watch Paths": "70AD47", "Taxonomy": "A5A5A5"}.get(s.title, "D9D9D9")
    wb.save(out)

    print(f"niches={len(NICHES)} entries={len(path_rows)} unique_films={len(films)} "
          f"watched={len(watched)} seen_in_guide={seen_count}")
    for w in warnings:
        print("WARN", w)
    return films, watched


if __name__ == "__main__":
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else str(Path(__file__).resolve().parent / "Film_Watch_Guide.xlsx")
    main(src, out)
